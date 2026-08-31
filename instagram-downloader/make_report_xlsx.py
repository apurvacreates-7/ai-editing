#!/usr/bin/env python3
"""Build Fussy_cat_analysis.xlsx — a 4-tab Excel of the latest analysis,
using the SAME routing and overrides as build_dashboard.py (v15).

Run in the Fussy cat folder:
    python3 make_report_xlsx.py "Fussy chat - all with media.csv"
Needs: python3 -m pip install --user --break-system-packages openpyxl (already installed)
"""
import csv, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import glob, sys

def find_roster():
    for p in sorted(glob.glob("*.csv")):
        try:
            h = open(p, encoding="utf-8").readline().lower()
        except Exception:
            continue
        if "link of video" in h and "s3_link" in h:
            return p
    return None

ROSTER = sys.argv[1] if len(sys.argv) > 1 else find_roster()
if not ROSTER or not os.path.exists(ROSTER):
    print('Roster CSV not found. Pass it: python3 make_report_xlsx.py "sheet.csv"')
    sys.exit(1)
ANALYSIS = "analysis_results.csv"
SCAN = "scan_results.csv"
OUT = "Fussy_cat_analysis.xlsx"

AI_FINDINGS = {
    1443: ("AI-generated video", "AI-generated (Veo/Gemini watermark)"),
    949:  ("AI / fake photo", "AI-generated / composite image"),
    494:  ("Stock image", "Looks like a stock photo, not their own cat"),
    798:  ("Stock image", "Looks like a stock photo, not their own cat"),
    448:  ("Stock image", "Looks like a stock photo, not their own cat"),
    614:  ("Stock image", "Looks like a stock photo, not their own cat"),
    186:  ("Graphic / poster", "A designed poster, not a real photo"),
    650:  ("Graphic / listing", "A 'for sale' graphic, not a real photo"),
}
CONFIRMED_CATS = {n.lower() for n in [
    "Aryan Raj", "APRAJIT PODDAR", "Casius cochikunnel", "Mohd Danish zafar",
    "Shariffs crazy", "ARJUN", "Supriya Raju", "Mohd zaid", "Rajat Biswas",
    "Faizal hossain", "Thirunavukkarasu", "PARIMAL KUMAR GAMIT", "Bhuwan Sharma",
    "Ajoy shil", "Alisha mansoor", "Anosh m", "Daler singh", "Doli das",
    "Sarathy", "Senti", "RIKUL SARMA", "Hayat", "Twinkle Dutta", "Ali Ahmad",
]}
T1_CONFIRMED_CATS = {266, 989}
FORCE_DISQUALIFY = {n.lower() for n in [
    "Sara Huma", "Biswarup", "Suraj kumar", "Abhijit Dasgupta", "Ahsan masood",
    "Akmal Bari", "Vishal Kumar Das", "Neha", "Sarika goes", "Sk Lasammad",
]}

def is_ig(u): return bool(re.search(r"instagram\.com/(p|reel|reels|tv|stories)/", u or ""))
def cat_yes(s): return (s or "").strip().upper().startswith("CAT")
def cat_class(s):
    s = (s or "").strip()
    if not s: return "unknown"
    if s.upper().startswith("CAT"): return "cat"
    if "saw:" in s.lower(): return "other"
    return "unclear"
def cat_conf(s):
    m = re.search(r"(\d+)\s*%", s or "")
    return int(m.group(1)) if m else -1
def t1_status(v):
    if v in ("SAME", "LIKELY SAME"): return "match"
    if v == "DIFFERENT": return "different"
    if v == "UNCERTAIN": return "unclear"
    return "pending"
def pending_reason(link, note):
    l = (link or "").lower(); n = (note or "").lower()
    if "/stories/" in l or "no results" in n or "story could not be found" in n:
        return "Invalid link", "Submitted a Story link — Stories expire after 24 h, so it can't be verified"
    return "Couldn't fetch", "Submitted a link, but the post couldn't be fetched (private or removed)"
def simval(d):
    s = str(d.get("similarity", "")).replace("%", "").strip()
    try: return int(s)
    except Exception: return -1

def load_by_idx(path):
    d = {}
    for r in csv.DictReader(open(path, newline="", encoding="utf-8")):
        try: d[int(r["idx"])] = r
        except Exception: pass
    return d

analysis = load_by_idx(ANALYSIS)
scan = load_by_idx(SCAN)
rows = list(csv.DictReader(open(ROSTER, newline="", encoding="utf-8")))
print(f"roster {len(rows)} | analysis {len(analysis)} | scan {len(scan)}")

t1, t2, t3, t4 = [], [], [], []
for i, r in enumerate(rows, start=1):
    name = (r.get("name") or "").strip()
    link = (r.get("link of video") or "").strip()
    video = (r.get("video_s3_link") or "").strip()
    photo = (r.get("photo_s3_link") or "").strip()
    media_link = video or photo
    has_media = bool(media_link)
    mtype = "video" if video else ("photo" if photo else "")
    ai = AI_FINDINGS.get(i)

    if name.lower() in FORCE_DISQUALIFY and (has_media or is_ig(link)):
        t3.append([i, name, "Disqualified", "Reviewed — not a valid entry", mtype, media_link])
        continue

    if is_ig(link):
        a = analysis.get(i, {})
        verdict = a.get("match_verdict", "") or "PENDING"
        cat_ig = cat_yes(a.get("instagram_has_cat"))
        cat_chat = cat_yes(a.get("chat_image_has_cat"))
        if i in T1_CONFIRMED_CATS:
            cat_ig = cat_chat = True
        has_cat = "Yes" if (cat_ig or cat_chat) else "No"
        status = t1_status(verdict)
        note = ""
        if status == "pending":
            lbl, why = pending_reason(link, a.get("note", ""))
            note = why
            t4.append([i, name, link, why, "yes" if media_link else "no", mtype])
        t1.append({"row": [i, name, verdict, a.get("similarity", ""), has_cat,
                           "Yes" if cat_chat else "No", "Yes" if cat_ig else "No",
                           a.get("ig_username", ""), a.get("likes", ""),
                           a.get("comments", ""), a.get("views", ""),
                           a.get("instagram_caption", ""), link, media_link, note],
                   "similarity": a.get("similarity", "")})
        continue

    sc = scan.get(i, {})
    klass = cat_class(sc.get("has_cat")) if sc else "unknown"
    conf = cat_conf(sc.get("has_cat")) if sc else -1
    real_link = bool(link) and not link.lower().startswith("choice-")
    confirmed = name.lower() in CONFIRMED_CATS

    if has_media and confirmed:
        t2.append({"row": [i, name, mtype, "Yes", "Confirmed by review", media_link],
                   "k": "cat", "conf": 100})
        continue
    if has_media and ai:
        t3.append([i, name, ai[0], ai[1], mtype, media_link])
        continue
    if real_link:
        ltype = ("They linked their profile, not a post"
                 if "instagram.com" in link else "Not a valid Instagram post link")
        t4.append([i, name, link, ltype, "yes" if media_link else "no", mtype])
        continue
    if has_media:
        t2.append({"row": [i, name, mtype,
                           "Yes" if klass == "cat" else "No",
                           sc.get("has_cat", ""), media_link],
                   "k": klass, "conf": conf})

t1.sort(key=simval, reverse=True)
t2.sort(key=lambda d: (0 if d["k"] == "cat" else 1, -d["conf"]))
print(f"t1 {len(t1)} | t2 {len(t2)} | t3 {len(t3)} | t4 {len(t4)}")

# ---------- analysis batch (rows up to the cutoff = first pass) ----------
BATCH_CUTOFF = 2399
def batch_label(i):
    return "Batch 1 (24 Aug 2025)" if i <= BATCH_CUTOFF else "Batch 2 (31 Aug 2025)"
for d in t1: d["row"].insert(2, batch_label(d["row"][0]))
for d in t2: d["row"].insert(2, batch_label(d["row"][0]))
for row in t3: row.insert(2, batch_label(row[0]))
for row in t4: row.insert(2, batch_label(row[0]))

# ---------- workbook ----------
wb = Workbook()
NAVY = "16265B"
hdr_fill = PatternFill("solid", fgColor=NAVY)
hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
base = Font(name="Arial", size=10)
bold = Font(name="Arial", size=10, bold=True)

def sheet(ws, headers, data, widths):
    ws.freeze_panes = "A2"
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r, row in enumerate(data, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = base
            cell.alignment = Alignment(vertical="top", wrap_text=(widths[c - 1] >= 40))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"

ws1 = wb.active; ws1.title = "1 Instagram submissions"
sheet(ws1,
      ["Row", "Name", "Analysis batch", "Verdict", "Similarity", "Has cat", "Cat in chat",
       "Cat in post", "IG username", "Likes", "Comments", "Views", "Caption",
       "Instagram link", "Chat media link", "Note"],
      [d["row"] for d in t1],
      [6, 22, 20, 14, 10, 9, 10, 10, 22, 8, 10, 8, 45, 45, 45, 45])

ws2 = wb.create_sheet("2 No Instagram link")
sheet(ws2, ["Row", "Name", "Analysis batch", "Media type", "Has cat", "Detector detail",
            "Chat media link"],
      [d["row"] for d in t2], [6, 22, 20, 11, 9, 26, 60])

ws3 = wb.create_sheet("3 Disqualified")
sheet(ws3, ["Row", "Name", "Analysis batch", "Category", "Reason", "Media type",
            "Chat media link"],
      t3, [6, 22, 20, 20, 45, 11, 60])

ws4 = wb.create_sheet("4 Invalid links")
sheet(ws4, ["Row", "Name", "Analysis batch", "Link submitted", "Why it's invalid",
            "Has chat media", "Media type"],
      t4, [6, 22, 20, 50, 55, 14, 11])

# ---------- summary ----------
wsS = wb.create_sheet("Summary", 0)
wsS.column_dimensions["A"].width = 46
wsS.column_dimensions["B"].width = 12
wsS["A1"] = "Fussy Cat UGC — analysis summary"; wsS["A1"].font = Font(name="Arial", size=13, bold=True)
labels = [
    ("Instagram submissions reviewed", "'1 Instagram submissions'"),
    ("No Instagram link (uploaded media only)", "'2 No Instagram link'"),
    ("Disqualified", "'3 Disqualified'"),
    ("Invalid links (incl. links that couldn't be fetched)", "'4 Invalid links'"),
]
for j, (lab, ref) in enumerate(labels, start=3):
    wsS.cell(row=j, column=1, value=lab).font = base
    c = wsS.cell(row=j, column=2, value=f"=COUNTA({ref}!A:A)-1")
    c.font = bold
wsS["A8"] = "Instagram tab — has cat"; wsS["A8"].font = base
wsS["B8"] = "=COUNTIF('1 Instagram submissions'!F:F,\"Yes\")"; wsS["B8"].font = bold
wsS["A9"] = "Instagram tab — does not have cat"; wsS["A9"].font = base
wsS["B9"] = "=COUNTIF('1 Instagram submissions'!F:F,\"No\")"; wsS["B9"].font = bold
wsS["A11"] = ("Note: 'Has cat' on the Instagram tab includes two human-confirmed rows (266, 989). "
              "A pending/not-fetched Instagram link is listed on tab 1 AND tab 4, matching the dashboard.")
wsS["A11"].font = Font(name="Arial", size=9, italic=True, color="666666")
wsS["A11"].alignment = Alignment(wrap_text=True)
wsS.merge_cells("A11:B13")

wb.save(OUT)
print("saved", OUT)
